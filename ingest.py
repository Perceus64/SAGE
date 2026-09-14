"""Turn course material files into plain-text units ready for chunking.

Supports PDFs, PowerPoint decks, and plain text / markdown notes. Each
extracted unit carries lightweight provenance (source file, page/slide
number) so answers can later be traced back to where they came from.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".pptx", ".txt", ".md"}


@dataclass
class RawUnit:
    """One extracted piece of text before chunking (e.g. one PDF page)."""

    source_path: str
    location: str  # e.g. "page 3", "slide 7", "whole file"
    text: str


def file_hash(path: Path) -> str:
    """Content hash used to detect whether a file changed since last sync."""
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()


def _ingest_pdf(path: Path) -> list[RawUnit]:
    from pypdf import PdfReader

    units: list[RawUnit] = []
    try:
        reader = PdfReader(str(path))
    except Exception as e:  # corrupt / encrypted / unreadable PDF
        print(f"  ! skipped {path.name}: could not open PDF ({e})")
        return units

    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = text.strip()
        if text:
            units.append(RawUnit(str(path), f"page {i + 1}", text))
    return units


def _ingest_pptx(path: Path) -> list[RawUnit]:
    from pptx import Presentation

    units: list[RawUnit] = []
    try:
        prs = Presentation(str(path))
    except Exception as e:
        print(f"  ! skipped {path.name}: could not open PPTX ({e})")
        return units

    for i, slide in enumerate(prs.slides):
        pieces = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    pieces.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            pieces.append(cell.text.strip())
        # speaker notes are often where the real explanation lives
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                pieces.append(f"[notes] {note}")
        text = "\n".join(pieces).strip()
        if text:
            units.append(RawUnit(str(path), f"slide {i + 1}", text))
    return units


def _ingest_text(path: Path) -> list[RawUnit]:
    try:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
    except Exception as e:
        print(f"  ! skipped {path.name}: could not read file ({e})")
        return []
    if not text:
        return []
    return [RawUnit(str(path), "whole file", text)]


def ingest_file(path: Path) -> list[RawUnit]:
    """Extract raw text units from a single supported file."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _ingest_pdf(path)
    if suffix == ".pptx":
        return _ingest_pptx(path)
    if suffix in (".txt", ".md"):
        return _ingest_text(path)
    return []


def discover_files(folder: Path) -> list[Path]:
    """Find all supported course-material files under a folder, recursively."""
    return sorted(
        p
        for p in folder.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    )
